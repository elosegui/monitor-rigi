# -*- coding: utf-8 -*-
import numpy as np, csv, json
from eval_iiep import run_rz, APROB, provs, SEC, RATE
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
r=run_rz(APROB); NP,NS=24,20
def g(x): return np.asarray(x).reshape(NP,NS)
Etot=g(r['Edir']+r['Eind']+r['Eindu']).sum(1); Ef=g(r['Ef']).sum(1); Einf=Etot-Ef
Ed=g(r['Edir']).sum(1); Ei=g(r['Eind']).sum(1); Eu=g(r['Eindu']).sum(1)
stock=np.zeros(NP); pi={p:i for i,p in enumerate(provs)}
for row in csv.DictReader(open('/sessions/eloquent-awesome-mayer/mnt/RIGI/replicacion/rep_meta.csv')):
    stock[pi[row['provincia']]]+=float(row['empleo'])
wb=openpyxl.Workbook()
B=Font(bold=True); BW=Font(bold=True,color="FFFFFF"); TIT=Font(bold=True,size=14)
H=PatternFill("solid",fgColor="1F6E43"); HL=PatternFill("solid",fgColor="D9EAD3"); TOTf=PatternFill("solid",fgColor="FCE4D6")
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
cen=Alignment(horizontal="center",vertical="center"); lft=Alignment(horizontal="left",vertical="center")
ws=wb.active; ws.title="Empleo por provincia"
ws["A1"]="RIGI - Empleo del primer ano por provincia (22 proyectos aprobados)"; ws["A1"].font=TIT
ws["A2"]="Base: inversion comprometida Ano 1 (IIEP, Monitor RIGI) = USD 5.511 M. Empleo directo+indirecto+inducido via MIP provincial 2023; formalidad ajustada por provincia."
ws["A2"].font=Font(italic=True,size=9)
cols=["Provincia","Empleo total","Formal","Informal","% formal","Directo","Indirecto","Inducido","% del empleo provincial"]
wid=[20,13,11,11,10,11,11,11,20]
row=4
for j,(c,w) in enumerate(zip(cols,wid),1):
    cc=ws.cell(row,j,c); cc.font=BW; cc.fill=H; cc.alignment=cen; cc.border=bd
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width=w
row=5
for k in np.argsort(-Etot):
    if Etot[k]<0.5: continue
    vals=[provs[k].replace('_',' '),round(Etot[k]),round(Ef[k]),round(Einf[k]),
          "%.0f%%"%(100*Ef[k]/Etot[k]),round(Ed[k]),round(Ei[k]),round(Eu[k]),
          "%.2f%%"%(100*Etot[k]/stock[k] if stock[k] else 0)]
    for j,v in enumerate(vals,1):
        c=ws.cell(row,j,v); c.border=bd; c.alignment=lft if j==1 else cen
        if j in(2,3,4,6,7,8) and isinstance(v,int): c.number_format='#,##0'
    row+=1
tot=[ "TOTAL",round(Etot.sum()),round(Ef.sum()),round(Einf.sum()),"%.0f%%"%(100*Ef.sum()/Etot.sum()),
      round(Ed.sum()),round(Ei.sum()),round(Eu.sum()),""]
for j,v in enumerate(tot,1):
    c=ws.cell(row,j,v); c.border=bd; c.font=B; c.fill=TOTf; c.alignment=lft if j==1 else cen
    if j in(2,3,4,6,7,8): c.number_format='#,##0'
ws.freeze_panes="A5"; ws.sheet_view.showGridLines=False
out="/sessions/eloquent-awesome-mayer/mnt/RIGI/RIGI_empleo_provincia_IIEP.xlsx"
try:
    wb.save(out); print("OK",out)
except PermissionError:
    out=out.replace(".xlsx","_v2.xlsx"); wb.save(out); print("OK(v2)",out)
